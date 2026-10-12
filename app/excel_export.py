from collections import defaultdict
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD=PatternFill('solid',fgColor='19324A')
def date(value):
    if not value:return ''
    try:return datetime.fromisoformat(value).astimezone().strftime('%d.%m.%Y %H:%M:%S')
    except Exception:return str(value)
def sheet(book,name,headers,rows):
    ws=book.create_sheet(name);ws.append(headers)
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=HEAD
    for row in rows: ws.append(row)
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for n in range(1,len(headers)+1): ws.column_dimensions[get_column_letter(n)].width=min(42,max(12,len(headers[n-1])+2))
    return ws
def workbook(storage):
    matches,goals,odds,events=storage.export_bundle(); goals_by=defaultdict(list); odds_by=defaultdict(list)
    for g in goals:goals_by[g['match_id']].append(g)
    for o in odds:odds_by[o['match_id']].append(o)
    book=Workbook();book.remove(book.active)
    sheet(book,'Матчи',['ID матча','Дата','Время начала','Команда 1','Команда 2','Первый тайм','Второй тайм','Итог','Победитель','Голов 1','Голов 2','Всего','Последовательность команд','Последовательность счёта','Статус','Полнота'],[
        [m['external_match_id'],date(m['created_at']),date(m['started_at']),m['team_1'],m['team_2'],m['first_half_score'],m['second_half_score'],m['final_score'],m['winner'],*(map(int,(m['final_score'] or '0:0').split(':'))),m['total_goals'],' → '.join(str(g['team'] or '?') for g in goals_by[m['id']]),' → '.join(g['score_after'] for g in goals_by[m['id']]),m['status'],'Полная' if m['data_complete'] else 'Неполная'] for m in matches])
    sheet(book,'Голы',['ID матча','Дата матча','Команда 1','Команда 2','№','Тайм','Время гола','Кто забил','Счёт до','Счёт после','КФ 1','КФ 2','КФ забившей','КФ другой','Время наблюдения','Достоверность'],[
        [g['external_match_id'],date(g['started_at']),g['team_1'],g['team_2'],g['goal_number'],g['half'],g['game_time'],g['team'],g['score_before'],g['score_after'],g['team_1_odds'],g['team_2_odds'],g['team_1_odds'] if g['team']==g['team_1'] else g['team_2_odds'],g['team_2_odds'] if g['team']==g['team_1'] else g['team_1_odds'],date(g['detected_at']),g['sequence_confidence']] for g in goals])
    sheet(book,'Коэффициенты',['ID матча','Дата','Команда 1','Команда 2','Счёт','№ ожидаемого гола','КФ 1','КФ 2','Статус рынка','Время наблюдения'],[[o['external_match_id'],date(o['started_at']),o['team_1'],o['team_2'],o['score'],o['next_goal_number'],o['team_1_odds'],o['team_2_odds'],o['market_status'],date(o['captured_at'])] for o in odds])
    stats=defaultdict(lambda:defaultdict(int))
    for m in matches:
        if not m['data_complete'] or not m['final_score']:continue
        a,b=map(int,m['final_score'].split(':'))
        for team,gf,ga in ((m['team_1'],a,b),(m['team_2'],b,a)):
            s=stats[team];s['matches']+=1;s['gf']+=gf;s['ga']+=ga;s['wins']+=gf>ga;s['draws']+=gf==ga;s['losses']+=gf<ga
    sheet(book,'Команды',['Команда','Матчей','Побед','Ничьих','Поражений','Забито','Пропущено','Среднее забито'],[[k,v['matches'],v['wins'],v['draws'],v['losses'],v['gf'],v['ga'],round(v['gf']/v['matches'],3)] for k,v in stats.items()])
    known=defaultdict(int); missed=defaultdict(int)
    for g in goals:known[g['match_id']]+=1
    for e in events:
        if e['event_type']=='MISSED_SCORE_INTERVAL':missed[e['match_id']]+=1
    sheet(book,'Качество данных',['ID','Дата','Команды','Полнота','Известных голов','Пропусков','Записей коэффициентов','1-й тайм','2-й тайм','Итог'],[[m['external_match_id'],date(m['created_at']),f"{m['team_1']} — {m['team_2']}",'Полная' if m['data_complete'] else 'Неполная',known[m['id']],missed[m['id']],len(odds_by[m['id']]),bool(m['first_half_score']),bool(m['second_half_score']),bool(m['final_score'])] for m in matches])
    out=BytesIO();book.save(out);return out.getvalue()
