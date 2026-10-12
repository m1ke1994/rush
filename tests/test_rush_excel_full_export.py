"""Regression: Excel contains every observation, including unfinished games."""
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from openpyxl import load_workbook

from app.excel_export import workbook
from app.storage import Storage


def test_full_excel_round_trip_keeps_every_observation():
    with TemporaryDirectory() as tmp:
        store = Storage(Path(tmp) / "observations.sqlite3")
        first = store.upsert_match("ID-1", "Rush", "Наполи", "Челси")
        second = store.upsert_match("ID-2", "Rush", "Наполи", "Челси")
        store.set_match(first, status="FINISHED", final_score="1:0", data_complete=1)
        store.set_match(second, status="INTERRUPTED", data_complete=0)
        store.odds(first, "0:0", 1.81, 2.01, "AVAILABLE", "CANVAS_2D")
        store.odds(first, "0:0", 1.83, 2.02, "AVAILABLE", "CANVAS_2D")
        store.odds(first, "0:0", None, None, "NOT_FOUND", "CANVAS_2D", "Missing")
        store.goal(first, 1, "Наполи", "0:0", "1:0", "1", "01:05", "HIGH")
        store.event(second, "MISSED_SCORE_INTERVAL", {"before": "0:0", "after": "2:1"})

        exported = workbook(store)
        with ZipFile(BytesIO(exported)) as zip_file:
            assert zip_file.testzip() is None
            assert "xl/workbook.xml" in zip_file.namelist()

        book = load_workbook(BytesIO(exported), read_only=True, data_only=True)
        try:
            assert book["Матчи_исходные"].max_row == 3
            assert book["Голы_исходные"].max_row == 2
            assert book["КФ_все_снимки"].max_row == 4
            assert book["События_все"].max_row == 2
            assert book["Матчи_для_анализа"].max_row == 3
            assert book["Таблица_расчётная"].max_row >= 2
            match_ids = [row[1] for row in list(book["Матчи_исходные"].values)[1:]]
            assert set(match_ids) == {"ID-1", "ID-2"}
            assert any(row[10] == "NOT_FOUND" for row in list(book["КФ_все_снимки"].values)[1:])
        finally:
            book.close()


def test_empty_export_is_still_valid_xlsx():
    with TemporaryDirectory() as tmp:
        store = Storage(Path(tmp) / "empty.sqlite3")
        book = load_workbook(BytesIO(workbook(store)), read_only=True)
        try:
            assert book["Матчи_исходные"].max_row == 1
            assert book["О выгрузке"]["B2"].value == 0
        finally:
            book.close()
